from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables" / "guides"


COMPANY = "佛山市思科达医疗器械有限公司 / CICADA Dental"
SERVICE_PHONE = "+86-0757-85775667"
SERVICE_EMAIL = "info@cicadadental.com"
REPAIR_RECEIVER = "姚兵，13929198537"
REPAIR_ADDRESS = "广东省佛山市南海区狮山镇罗村广东新光源核心基地B5座五楼"


def set_run_font(run, name="Microsoft YaHei", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
        name,
    )
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", style=None, size=11, bold=False, color=None, align=None):
    paragraph = doc.add_paragraph(style=style)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return paragraph


def add_title(doc, title, subtitle):
    title_p = add_para(doc, title, size=22, bold=True, color="0F1F3A")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_p = add_para(doc, subtitle, size=11, color="5A6C8D")
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, f"适用对象：使用 {COMPANY} 售后小程序的诊所、医院及经销客户", size=10, color="5A6C8D", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "")


def add_h1(doc, text):
    p = doc.add_heading(level=1)
    run = p.add_run(text)
    set_run_font(run, size=16, bold=True, color="1E6FE0")
    return p


def add_h2(doc, text):
    p = doc.add_heading(level=2)
    run = p.add_run(text)
    set_run_font(run, size=13, bold=True, color="0F1F3A")
    return p


def add_bullet(doc, text):
    return add_para(doc, text, style="List Bullet")


def add_number(doc, text):
    return add_para(doc, text, style="List Number")


def add_meta(doc):
    add_h1(doc, "联系方式")
    add_bullet(doc, f"服务电话：{SERVICE_PHONE}")
    add_bullet(doc, f"服务邮箱：{SERVICE_EMAIL}")
    add_bullet(doc, f"寄修收件信息：{REPAIR_RECEIVER}，{REPAIR_ADDRESS}")


def add_source_note(doc):
    add_h1(doc, "资料依据")
    add_bullet(doc, "公司公开信息参考 CICADA Dental 官网、产品页和联系我们页面。")
    add_bullet(doc, "小程序操作说明依据当前项目中的首页模块、报修表单、查询页面、发票申请流程和后台操作教程上传功能整理。")
    add_bullet(doc, "开票税号、开户行、银行账号等财务信息未在公开资料中核验，本文档不作填写；以财务审核结果为准。")


def build_quick():
    doc = Document()
    add_title(doc, "快速指南", "3 分钟了解小程序售后服务入口")
    add_h1(doc, "可以办理什么")
    add_para(doc, "小程序面向牙科设备售后维修场景，支持提交报修、查看维修进度、查询包裹签收和处理状态、进行故障自查、申请电子发票，并查看保修政策、收费说明和客服联系信息。")
    add_h1(doc, "推荐使用顺序")
    add_number(doc, "登录微信账号，进入首页。登录后可以保存并查看与当前账号关联的维修工单。")
    add_number(doc, "需要寄修时，点击“立即报修”，填写产品、故障、寄出物流和回寄地址。")
    add_number(doc, "已寄出设备后，可用“包裹查询”输入快递单号，查看是否签收、登记和关联工单。")
    add_number(doc, "维修处理中，可在“维修进度”查看提交、运输、签收、处理、回寄和完成状态。")
    add_number(doc, "维修完成并结算后，可在“发票与开票”选择对应工单提交开票申请。")
    add_h1(doc, "首页主要入口")
    add_bullet(doc, "基础服务：立即报修、维修进度、包裹查询。")
    add_bullet(doc, "自助查询：故障自查、保修政策、收费指南。")
    add_bullet(doc, "操作教程：快速指南、报修指南、查询指南、开票指南。")
    add_bullet(doc, "我的：维修订单、收货地址、投诉建议、发票与开票。")
    add_h1(doc, "使用提醒")
    add_bullet(doc, "寄修前请备份并确认设备附件，随包裹附上故障描述、联系人和回寄地址。")
    add_bullet(doc, "工单、支付和发票功能与当前微信账号关联，建议用同一微信账号办理完整流程。")
    add_bullet(doc, "如页面显示“未查询到”，请确认单号是否正确，或等待工作人员签收录入后再查询。")
    add_meta(doc)
    add_source_note(doc)
    return doc


def build_repair():
    doc = Document()
    add_title(doc, "报修指南", "设备寄修、资料填写和进度跟踪说明")
    add_h1(doc, "报修前准备")
    add_bullet(doc, "确认产品名称、型号、序列号和购买日期；如有购买凭证，可在表单中上传。")
    add_bullet(doc, "拍摄故障图片或视频，最多按小程序提示上传 3 个附件，便于工程师判断。")
    add_bullet(doc, "妥善包装设备，易损部件建议单独固定，避免运输二次损坏。")
    add_bullet(doc, "在快递内留纸条，写明故障现象、联系人、联系电话和回寄地址。")
    add_h1(doc, "小程序提交流程")
    add_number(doc, "进入首页，点击“立即报修”。")
    add_number(doc, "填写产品信息：产品名称、序列号、购买日期、购买凭证、故障描述、图片或视频。")
    add_number(doc, "填写寄出信息：选择物流公司，录入或扫码填写运单号。")
    add_number(doc, "填写回寄信息：收货人、手机号、详细地址和单位名称。")
    add_number(doc, "确认无误后提交。提交成功后，小程序会生成工单并可进入“维修进度”查看。")
    add_h1(doc, "寄修收件信息")
    add_bullet(doc, f"收件公司：{COMPANY}")
    add_bullet(doc, f"收件人及电话：{REPAIR_RECEIVER}")
    add_bullet(doc, f"收件地址：{REPAIR_ADDRESS}")
    add_h1(doc, "维修处理说明")
    add_bullet(doc, "工作人员签收后会登记包裹并关联工单，随后进入检测和处理流程。")
    add_bullet(doc, "保修期内且符合保修条件的设备按保修政策处理；保修期外或非保修范围会先报价，客户确认后再维修。")
    add_bullet(doc, "维修完成后按回寄信息寄回，并在工单中更新物流状态。")
    add_h1(doc, "常见注意事项")
    add_bullet(doc, "请勿只寄设备而不提交报修单，否则可能影响签收后的识别和处理速度。")
    add_bullet(doc, "多个产品可在同一报修单中继续增加产品条目，但每个产品请单独写清故障现象。")
    add_bullet(doc, "如回寄地址变化，请在处理完成前及时联系客服或在地址管理中维护最新信息。")
    add_meta(doc)
    add_source_note(doc)
    return doc


def build_query():
    doc = Document()
    add_title(doc, "查询指南", "工单、包裹和故障知识查询说明")
    add_h1(doc, "可以查询哪些内容")
    add_bullet(doc, "工单查询：查看当前微信账号提交过的维修记录、设备信息、费用状态和处理状态。")
    add_bullet(doc, "包裹查询：输入快递单号，查看待签收、已签收、已登记、处理中、已关联等节点。")
    add_bullet(doc, "故障知识：按设备型号、产品类型或故障关键词查看自查问题、排查建议和处理建议。")
    add_h1(doc, "工单进度查询")
    add_number(doc, "进入首页点击“维修进度”，或在“我的”中进入维修订单。")
    add_number(doc, "可按全部、已提交、运输中、已签收、处理中、已回寄、已完成等状态筛选。")
    add_number(doc, "点击工单卡片可查看时间线、报价、付款、回寄和发票状态。")
    add_h1(doc, "包裹查询")
    add_number(doc, "进入首页点击“包裹查询”。")
    add_number(doc, "输入、扫码或粘贴快递单号。")
    add_number(doc, "如需查看更完整轨迹，可填写收件人手机号后四位进行隐私校验。")
    add_number(doc, "若暂未查询到，请确认单号正确，或等待签收录入后再查询。")
    add_h1(doc, "故障自查")
    add_bullet(doc, "可输入设备型号、故障现象或关键词，例如根管马达、光固化灯、牙科手机等。")
    add_bullet(doc, "自查结果仅作为初步排查参考，若涉及安全、性能异常或无法复现的问题，建议提交报修。")
    add_h1(doc, "查询不到时如何处理")
    add_bullet(doc, "工单查询需要登录，且只展示当前微信账号提交的维修记录。")
    add_bullet(doc, "包裹刚寄出时，可能还未签收或未录入系统，请稍后再查。")
    add_bullet(doc, "关键词过窄时可换用产品类型、故障现象或设备型号重新搜索。")
    add_meta(doc)
    add_source_note(doc)
    return doc


def build_invoice():
    doc = Document()
    add_title(doc, "开票指南", "发票申请、资料填写和开票进度查看说明")
    add_h1(doc, "适用条件")
    add_bullet(doc, "维修工单完成并已完成结算后，可在小程序中提交开票申请。")
    add_bullet(doc, "当前小程序支持电子普通发票申请；具体开票类型和开具结果以财务审核为准。")
    add_bullet(doc, "发票金额以已结算的维修费用或后台核准金额为准。")
    add_h1(doc, "申请路径")
    add_number(doc, "进入“我的”，点击“发票与开票”；也可从维修订单详情进入开票入口。")
    add_number(doc, "在“待开票”列表中选择需要开票的已完成工单。")
    add_number(doc, "点击“申请开票”，填写发票资料并确认提交。")
    add_number(doc, "提交后可在列表中查看审核中、开票中、已开票等状态。")
    add_h1(doc, "需要填写的资料")
    add_bullet(doc, "发票类型：默认电子普通发票。")
    add_bullet(doc, "抬头类型：企业单位或个人。企业单位适合诊所、医院和公司主体；个人无需填写税号。")
    add_bullet(doc, "发票抬头：请按营业执照或收票主体准确填写。")
    add_bullet(doc, "税号：企业单位必填纳税人识别号；个人抬头不需要填写。")
    add_bullet(doc, "接收邮箱：用于接收电子发票链接或文件，必须填写有效邮箱。")
    add_bullet(doc, "备注：如有特殊开票说明，可在备注中填写，最终以财务审核为准。")
    add_h1(doc, "开票状态说明")
    add_bullet(doc, "待申请：工单已满足申请条件，但尚未提交资料。")
    add_bullet(doc, "审核中：客服或财务正在核对抬头、税号、邮箱和金额。")
    add_bullet(doc, "开票中：资料已通过，等待财务开具电子发票。")
    add_bullet(doc, "已开票：电子发票已开具，可复制链接或按页面提示查看。")
    add_h1(doc, "重要说明")
    add_bullet(doc, "请确保抬头、税号和邮箱准确无误；资料错误可能导致退回或重新开具。")
    add_bullet(doc, "公开渠道未核验到公司专门开票指南或收款账户信息，本文档不填写税号、开户行和银行账号。")
    add_bullet(doc, "如需专票、红冲、重开或合并开票，请通过客服或后台财务人员确认后处理。")
    add_meta(doc)
    add_source_note(doc)
    return doc


def apply_document_defaults(doc):
    section = doc.sections[0]
    section.top_margin = Pt(72)
    section.bottom_margin = Pt(72)
    section.left_margin = Pt(72)
    section.right_margin = Pt(72)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
        "Microsoft YaHei",
    )
    normal.font.size = Pt(11)
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.2


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    docs = {
        "快速指南.docx": build_quick(),
        "报修指南.docx": build_repair(),
        "查询指南.docx": build_query(),
        "开票指南.docx": build_invoice(),
    }
    for name, doc in docs.items():
        apply_document_defaults(doc)
        doc.save(OUT_DIR / name)
        print(OUT_DIR / name)


if __name__ == "__main__":
    main()
